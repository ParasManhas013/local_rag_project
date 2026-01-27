from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import json
import uuid
import os

TEMPLATE_PATH = "Navy Blue and Purple Gradient Futuristic Computer Technology Presentation.pptx"

# Mapping structure based on 'inspect_slides.py' output
# Keys are Slide Indices (0-based)
# Values are dictionaries mapping "logical_role" -> [Shape IDs to update]
# "clear" is a list of Shape IDs to empty out (e.g. split titles like "Hardware" "Components")

SLIDE_MAPPING = {
    0: { # Title Slide
        "title": [3],       # "Computer" -> Main Title
        "subtitle": [5],    # "Lorem ipsum..." -> Subtitle
        "clear": [4]        # "Technology" -> Clear it (assume Title covers it or fits in 3)
    },
    1: { # Introduction
        "title": [14],      # "Introduction"
        "content": [15]     # Main Body Text
    },
    2: { # Hardware (2 Text Blocks)
        "title": [14],      # "Hardware"
        "content": [18, 19],# "Lorem..." (Top), "Lorem..." (Bottom) -> Split content here
        "clear": [15, 16, 17] # "Components", "CPU", "Memory" -> Clear sub-headers
    },
    3: { # Software
        "title": [16],      # "Software"
        "content": [14],    # Main Body Text
        "clear": [17]       # "Components"
    },
    4: { # Networking (2 Text Blocks)
        "title": [14],      # "Networking"
        "content": [16, 17],# Two body text blocks
        "clear": [15]       # "& the Internet"
    },
    5: { # Data Management (3 Columns)
        "title": [16],      # "Data"
        "content": [19, 23], # Use Left and Right columns only to avoid overlap
        "clear": [17, 18, 20, 21, 22] # Clear Management, all headers, and middle body
    },
    6: { # Emerging Tech (2 Columns)
        "title": [14],      # "Emerging"
        "content": [16, 17],
        "clear": [15]       # "Technologies"
    },
    7: { # Cybersecurity (3 Columns)
        "title": [15],
        "content": [16, 17, 18],
        "clear": []
    },
    8: { # Impact (1 Text Block)
        "title": [14],
        "content": [18],
        "clear": [15]
    }
}

def apply_formatting(paragraph, font_size_pt=None, is_title=False):
    """
    Applies custom formatting: White Color, Arial font, Bold.
    Prioritizes font_size_pt if provided. Otherwise uses 36 for titles, 19 for body.
    """
    paragraph.font.name = 'Arial'
    paragraph.font.color.rgb = RGBColor(255, 255, 255)
    paragraph.font.bold = True
    
    if font_size_pt is not None:
        paragraph.font.size = Pt(font_size_pt)
    elif is_title:
        paragraph.font.size = Pt(36)
    else:
        paragraph.font.size = Pt(19)

def create_ppt_from_json(json_string: str):
    try:
        data = json.loads(json_string)
    except json.JSONDecodeError:
        print(f"Error: Invalid JSON string provided to create_ppt_from_json: {json_string}")
        return None

    # Load the template
    if not os.path.exists(TEMPLATE_PATH):
        print(f"Error: Template file not found at {TEMPLATE_PATH}")
        return None
        
    try:
        prs = Presentation(TEMPLATE_PATH)
    except Exception as e:
        print(f"Error loading template: {e}")
        return None

    # We do NOT delete slides. We modify existing ones.
    # We iterate through the JSON slides and map them to the Template slides by index.
    
    slides_data = data.get("slides", [])
    
    # Iterate strictly through the available mapping to avoid errors
    for i, mapping in SLIDE_MAPPING.items():
        if i >= len(prs.slides):
            break # No more slides in template
            
        if i >= len(slides_data):
            # We have more template slides than content. 
            # We could clear them or leave them. Let's leave them for now or clear key text.
            continue
            
        slide_content = slides_data[i]
        slide = prs.slides[i]
        
        # 1. Clear specified shapes
        for clear_id in mapping.get("clear", []):
            for shape in slide.shapes:
                if shape.shape_id == clear_id and shape.has_text_frame:
                    shape.text_frame.text = ""

        # 2. Update Title
        title_ids = mapping.get("title", [])
        new_title = slide_content.get("title", "")
        if title_ids and new_title:
            target_id = title_ids[0]
            for shape in slide.shapes:
                if shape.shape_id == target_id and shape.has_text_frame:
                    # Clear existing and add new paragraph to ensure formatting applies cleanly
                    tf = shape.text_frame
                    tf.clear()
                    p = tf.add_paragraph()
                    p.text = new_title
                    
                    if i == 0:
                        # Page 1 Title: Larger Font
                        apply_formatting(p, font_size_pt=54)
                    elif i == 5:
                        # Page 6 Title: Smaller Font to prevent breaking
                        apply_formatting(p, font_size_pt=28)
                    else:
                        apply_formatting(p, is_title=True)

        # 3. Update Subtitle (Slide 0 only usually)
        subtitle_ids = mapping.get("subtitle", [])
        new_subtitle = slide_content.get("subtitle", "")
        if subtitle_ids and new_subtitle:
             target_id = subtitle_ids[0]
             for shape in slide.shapes:
                if shape.shape_id == target_id and shape.has_text_frame:
                    tf = shape.text_frame
                    tf.clear()
                    p = tf.add_paragraph()
                    p.text = new_subtitle
                    apply_formatting(p, font_size_pt=24) # Slightly larger for subtitle

        # 4. Update Content
        content_ids = mapping.get("content", [])
        new_content = slide_content.get("content", []) # Expecting a list of strings
        
        # Convert string to list if necessary
        if isinstance(new_content, str):
            new_content = [new_content]
            
        if content_ids and new_content:
            # Strategy: Distribute content points across the available content text boxes.
            num_boxes = len(content_ids)
            points_per_box = (len(new_content) + num_boxes - 1) // num_boxes
            
            for idx, shape_id in enumerate(content_ids):
                # Find the shape
                target_shape = None
                for s in slide.shapes:
                    if s.shape_id == shape_id:
                        target_shape = s
                        break
                
                if target_shape and target_shape.has_text_frame:
                    tf = target_shape.text_frame
                    tf.clear() # Clear "Lorem Ipsum"
                    
                    # Get slice of content for this box
                    start = idx * points_per_box
                    end = start + points_per_box
                    chunk = new_content[start:end]
                    
                    for point in chunk:
                        p = tf.add_paragraph()
                        # Clean the point text first
                        clean_point = point.strip()
                        # Ensure it ends with a period
                        if not clean_point.endswith('.'):
                            clean_point += "."
                        # Format as "- Text."
                        p.text = f"- {clean_point}"
                        p.level = 0
                        apply_formatting(p, font_size_pt=19)

    # --- Save the presentation ---
    filename = f"{uuid.uuid4()}.pptx"
    output_dir = os.path.join("storage", "generated")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    prs.save(output_path)
    return output_path
